#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""合并MEE深挖结果+PBC爬取 -> 最终Excel"""

import json, os, re, time, random, urllib3
import requests
from bs4 import BeautifulSoup
from collections import Counter
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
urllib3.disable_warnings()

h = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
GREEN = ["绿色金融","绿色信贷","ESG","碳排放","碳中和","碳金融","绿色债券",
         "环境信息","节能减排","低碳","清洁能源","新能源","碳达峰","气候",
         "可持续发展","温室气体","碳汇","污染防治","环保","生态",
         "转型金融","赤道原则","气候投融资","社会责任","公司治理",
         "碳交易","排污权","能效","环境社会治理"]

def http_get(url):
    for _ in range(2):
        try:
            r = requests.get(url, headers=h, timeout=15, allow_redirects=True)
            if r.status_code == 200:
                if r.encoding and r.encoding.upper() in ("ISO-8859-1","LATIN-1"):
                    r.encoding = "utf-8"
                if not re.search(r"[一-鿿]", r.text[:5000]):
                    r.encoding = "gbk"
                return r
            if r.status_code == 404: return r
        except: time.sleep(2)
    return None

def extract_date(text, url=""):
    m = re.search(r"(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})", text)
    if m: return f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"
    m = re.search(r"/(\d{4})(\d{2})(\d{2})", url)
    if m: return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    return ""

def extract_doc(text):
    for p in [r"([一-鿿]+〔\d{4}〕\d+号)",
              r"((?:银发|银监发|证监发|银保监发|环发|国发|国办发)〔\d{4}〕\d+号)"]:
        m = re.search(p, text)
        if m: return m.group(1)
    return ""

# ===== 1. Load MEE deep results =====
print("Loading MEE deep results...")
with open("output/mee_results.json", "r", encoding="utf-8") as f:
    mee_results = json.load(f)
print(f"  MEE: {len(mee_results)}")

# ===== 2. PBC crawl =====
print("Running PBC crawl...")
pbc_base = "https://www.pbc.gov.cn"
pbc_channels = [
    ("/goutongjiaoliu/113456/113469/index.html",
     r'href="(/goutongjiaoliu/\d+/\d+/\d{14,}/index\.html)"', "新闻发布"),
    ("/tiaofasi/144941/144957/index.html",
     r'href="(/tiaofasi/\d+/\d+/\d{14,}/index\.html)"', "金融法规"),
    ("/zhengcehuobisi/125207/125213/index.html",
     r'href="(/\w+/\d+/\d+/\d{6,}/index\.html)"', "货币政策公告"),
    ("/zhengcehuobisi/125207/125213/125439/index.html",
     r'href="(/\w+/\d+/\d+/\d+/\d{6,}/index\.html)"', "信贷政策"),
    ("/zhengcehuobisi/125207/125213/125435/index.html",
     r'href="(/\w+/\d+/\d+/\d+/\d{6,}/index\.html)"', "公开市场"),
    ("/zhengcehuobisi/125207/125213/125431/index.html",
     r'href="(/\w+/\d+/\d+/\d+/\d{6,}/index\.html)"', "货币政策委员会"),
    ("/yanjiuju/3911332/index.html",
     r'href="(/\w+/\d+/\d+/\d{6,}/index\.html)"', "研究局"),
    ("/jinrongshichangsi/147160/index.html",
     r'href="(/jinrongshichangsi/\d+/\d+/\d{6,}/index\.html)"', "金融市场司"),
    ("/jinrongwendingju/146766/index.html",
     r'href="(/\w+/\d+/\d+/\d{6,}/index\.html)"', "金融稳定局"),
    ("/diaochatongjisi/116219/index.html",
     r'href="(/\w+/\d+/\d+/\d{6,}/index\.html)"', "调查统计司"),
    ("/xindaishichangsi/5443861/index.html",
     r'href="(/\w+/\d+/\d+/\d{6,}/index\.html)"', "信贷市场司"),
    ("/kejisi/146812/index.html",
     r'href="(/\w+/\d+/\d+/\d{6,}/index\.html)"', "科技司"),
    ("/huobizhengceersi/214481/index.html",
     r'href="(/\w+/\d+/\d+/\d{6,}/index\.html)"', "货币政策二司"),
]

seen_urls = set()
pbc_results = []

for ch, regex, cat in pbc_channels:
    r = http_get(pbc_base + ch)
    if not r or r.status_code != 200: continue

    hrefs = set(re.findall(regex, r.text))
    for href in hrefs:
        if href in seen_urls: continue
        seen_urls.add(href)

        full = "https://www.pbc.gov.cn" + href
        ar = http_get(full)
        if not ar or ar.status_code != 200 or len(ar.text) < 500: continue

        soup = BeautifulSoup(ar.text, "lxml")
        body = soup.get_text()

        title = ""
        for s in ["h1","h2",".article-title",".xl-title",".con-title","title"]:
            e = soup.select_one(s)
            if e:
                t = e.get_text(strip=True)
                if len(t) > 5: title = re.sub(r"[-–—|_].*$","",t).strip(); break

        if not title or len(title) < 5: continue
        if any(bl in title for bl in ["政府网站年度","网站地图","APP下载","版权声明"]): continue

        summary = ""
        for s in ["div.article-content","div.content","div.detail-content","div.article","div.con"]:
            e = soup.select_one(s)
            if e: summary = re.sub(r"<[^>]+>"," ",str(e))[:800]; break
        summary = re.sub(r"\s+"," ",summary).strip()

        if any(kw in (title + " " + summary + " " + body) for kw in GREEN):
            pbc_results.append({
                "title": title, "issuer": "中国人民银行",
                "doc_number": extract_doc(body),
                "pub_date": extract_date(body, href),
                "category": cat, "url": full, "summary": summary
            })
            print(f"  [PBC-{cat}] {title[:70]}")

    time.sleep(random.uniform(0.2, 0.5))

print(f"  PBC: {len(pbc_results)}")

# ===== 3. Combine + Dedup =====
all_results = mee_results + pbc_results
seen_u = set()
unique = []
for r in all_results:
    u = r["url"].strip().lower()
    if u and u not in seen_u:
        seen_u.add(u)
        unique.append(r)

print(f"\nTotal unique: {len(unique)}")

# ===== 4. Excel =====
fp = "output/policies_green_finance.xlsx"
os.makedirs("output", exist_ok=True)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "绿色金融政策文件"

hdrs = ["序号","文件名称","发文机构","文号","发布时间","政策类别","原文链接","摘要/主要内容"]
HF = PatternFill(start_color="2F5496",end_color="2F5496",fill_type="solid")
HFT = Font(name="微软雅黑",size=11,bold=True,color="FFFFFF")
CF = Font(name="微软雅黑",size=10)
BD = Border(left=Side(style="thin"),right=Side(style="thin"),
            top=Side(style="thin"),bottom=Side(style="thin"))

for ci, hdr in enumerate(hdrs, 1):
    c = ws.cell(row=1, column=ci, value=hdr)
    c.fill=HF; c.font=HFT; c.border=BD
    c.alignment=Alignment(horizontal="center",vertical="center")

for ri, rec in enumerate(unique, 2):
    for k, d in [("title","标题未收录"),("issuer","发文机构未收录"),
                 ("doc_number","文号未收录"),("pub_date","日期未收录"),
                 ("category","政策文件"),("url","链接未收录"),("summary","摘要未收录")]:
        if not rec.get(k,""): rec[k] = d

    row = [ri-1, rec["title"], rec["issuer"], rec["doc_number"],
           rec["pub_date"], rec["category"], rec["url"], rec["summary"]]
    for ci, v in enumerate(row, 1):
        c = ws.cell(row=ri, column=ci, value=v)
        c.font=CF; c.border=BD
        c.alignment = (Alignment(horizontal="center",vertical="top")
                      if ci in (1,4,5,6)
                      else Alignment(wrap_text=True,vertical="top"))

for ci, w in enumerate([6,55,20,24,12,16,55,75], 1):
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.freeze_panes = "A2"

wb.save(fp)
print(f"\nExcel: {fp} ({len(unique)} rows)")

# Stats
print("\nSource distribution:")
for issuer, cnt in Counter(r["issuer"] for r in unique).most_common():
    print(f"  {issuer}: {cnt}")
print(f"  TOTAL: {len(unique)}")

for attr, label in [("title","标题"),("pub_date","日期"),("summary","摘要")]:
    ok = sum(1 for r in unique if "未收录" not in r.get(attr,""))
    print(f"  {label}: {ok}/{len(unique)}")
print("Done!")
