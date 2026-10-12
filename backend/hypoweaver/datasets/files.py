from pathlib import Path
import hashlib
import numpy as np
import pandas as pd
import pyreadstat
from openpyxl import load_workbook
from ..errors import AppError

MAX_FILE = 500 * 1024**2
MAX_ROWS = 2_000_000
MAX_COLUMNS = 300
MAX_STORAGE = 3 * 1024**3


def inspect(path, sheet=None):
    path = Path(path)
    suffix = path.suffix.lower()
    labels, sheets = {}, []
    try:
        if suffix == '.dta':
            _, meta = pyreadstat.read_dta(path, metadataonly=True)
            columns = meta.column_names
            labels = dict(zip(columns, meta.column_labels))
        elif suffix == '.csv':
            columns = pd.read_csv(path, nrows=0).columns.tolist()
        elif suffix == '.xlsx':
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                sheets = book.sheetnames
                sheet = sheet or sheets[0]
                if sheet not in sheets:
                    raise AppError('sheet_not_found', '工作表不存在')
                values = next(book[sheet].iter_rows(values_only=True), ())
                columns = [str(x) if x is not None else f'Unnamed: {i}' for i, x in enumerate(values)]
                if len(columns) != len(set(columns)):
                    raise AppError('duplicate_columns', '列名存在重复，请在原始工作表中修正后导入')
            finally:
                book.close()
        else:
            raise AppError('unsupported_file', '仅支持 .csv、.xlsx 和 .dta')
    except AppError:
        raise
    except Exception as exc:
        raise AppError('inspect_failed', f'无法读取文件元数据（{type(exc).__name__}），请检查文件格式或编码') from exc
    if not columns:
        raise AppError('empty_file', '文件中没有可导入的列')
    return {'columns': [{'name': c, 'label': labels.get(c) or ''} for c in columns], 'sheets': sheets, 'sheet': sheet, 'format': suffix[1:]}


def read_selected(path, columns, sheet=None):
    path = Path(path)
    if not columns or len(columns) > MAX_COLUMNS or len(set(columns)) != len(columns):
        raise AppError('invalid_columns', '请选择 1–300 个不重复字段')
    try:
        if path.suffix.lower() == '.dta':
            frame, _ = pyreadstat.read_dta(path, usecols=columns, apply_value_formats=False, row_limit=MAX_ROWS + 1)
        elif path.suffix.lower() == '.xlsx':
            frame = pd.read_excel(path, sheet_name=sheet or 0, usecols=columns, nrows=MAX_ROWS + 1)
        else:
            frame = pd.read_csv(path, usecols=columns, nrows=MAX_ROWS + 1)
    except Exception as exc:
        raise AppError('import_failed', f'所选字段读取失败（{type(exc).__name__}），请检查列名、工作表和格式') from exc
    if len(frame) > MAX_ROWS or frame.empty:
        raise AppError('invalid_rows', '数据为空或超过 200 万行上限，请选择较小的数据文件')
    return frame[columns]


def profile(frame, labels=None):
    labels = labels or {}
    result = []
    for col in frame:
        series = frame[col]
        finite = pd.to_numeric(series, errors='coerce').replace([np.inf, -np.inf], np.nan)
        numeric = pd.api.types.is_numeric_dtype(series)
        result.append({'name': str(col), 'label': labels.get(col, ''), 'type': str(series.dtype),
                       'missing_rate': float(series.isna().mean()), 'unique_count': int(series.nunique()),
                       'min': float(finite.min()) if numeric and finite.notna().any() else None,
                       'max': float(finite.max()) if numeric and finite.notna().any() else None})
    return result


def panel_info(frame, unit=None, time=None):
    lookup = {str(c).lower(): c for c in frame}
    unit = unit or next((lookup[c] for c in ['firm_id', 'unit', 'symbol', 'company_id', 'id'] if c in lookup), None)
    time = time or next((lookup[c] for c in ['year', 'time', 'period'] if c in lookup), None)
    if not unit or not time:
        return {'unit': unit, 'time': time, 'detected': False}
    if unit not in frame or time not in frame or unit == time:
        raise AppError('invalid_panel', '个体和时间须选择两个不同的已导入字段')
    n_units, n_periods = int(frame[unit].nunique()), int(frame[time].nunique())
    duplicates = int(frame.duplicated([unit, time]).sum())
    missing_keys = int(frame[[unit, time]].isna().any(axis=1).sum())
    return {'unit': unit, 'time': time, 'detected': True, 'n_units': n_units, 'n_periods': n_periods,
            'duplicates': duplicates, 'missing_keys': missing_keys,
            'balanced': len(frame) == n_units*n_periods and duplicates == 0 and missing_keys == 0,
            'time_min': str(frame[time].min()), 'time_max': str(frame[time].max())}


def sha256(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as stream:
        for chunk in iter(lambda: stream.read(1024**2), b''):
            digest.update(chunk)
    return digest.hexdigest()


def check_storage(settings, incoming=0):
    root = settings.storage_dir()
    total = 0
    for name in ['staging', 'datasets', 'raw']:
        directory = root / name
        if directory.exists():
            for p in directory.iterdir():
                if p.is_file():
                    total += p.stat().st_size
    if total + incoming > MAX_STORAGE:
        raise AppError('storage_limit', '数据存储已达 3 GB 上限。可复用已上传文件，或按 README 整理不再使用的文件；系统不会自动删除历史文件')
